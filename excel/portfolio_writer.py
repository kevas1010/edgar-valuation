"""Portfolio workbook: one row per holding plus a History sheet that tracks the portfolio across runs.

Fidelity's numbers (value, cost basis) and each holding's model outputs are inputs (blue);
gains, weights, upside and totals are live formulas. Each modeled row links to that
holding's full valuation workbook.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.defined_name import DefinedName

from excel.styles import (
    FMT_PCT, FMT_PER_SHARE, GREY, KEY_FILL, THIN_TOP, calc, font, header_row, inp, label, title, widths,
)
from portfolio.run import MODELED, HoldingResult

FMT_USD = '$#,##0.00;($#,##0.00);"–"'
FMT_QTY = "#,##0.000"


def qty_format(q: float) -> str:
    """Whole shares as integers; fractional (fund) shares to 3 decimals. ("#,##0.###" shows "50." in Excel.)"""
    return "#,##0" if float(q).is_integer() else FMT_QTY
COLUMNS = ["Symbol", "Description", "Status", "Quantity", "Price", "Market value", "Cost basis",
           "Gain $", "Gain %", "Weight", "DCF value", "Comps value", "Model target", "Model upside",
           "Value × upside", "Value w/ target", "Workbook", "Notes"]


def write_portfolio_workbook(results: list[HoldingResult], history: pd.DataFrame, totals: pd.DataFrame,
                             path: Path, as_of: date) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    ws.sheet_view.showGridLines = False
    title(ws, "A1", "Portfolio – gain/loss and model valuation",
          f"Fidelity positions as exported · models run {as_of:%B %d, %Y} · model target = scenario-weighted "
          "DCF value per share · a model output, not investment advice")
    header_row(ws, 4, 1, COLUMNS)

    first = 5
    last = first + len(results) - 1
    for i, r in enumerate(results):
        row = first + i
        h = r.holding
        label(ws, f"A{row}", h.symbol, bold=True)
        label(ws, f"B{row}", h.description)
        label(ws, f"C{row}", r.status, color="000000" if r.status == MODELED else GREY)
        inp(ws, f"D{row}", h.quantity, qty_format(h.quantity))
        inp(ws, f"E{row}", h.last_price, FMT_USD)
        inp(ws, f"F{row}", h.current_value, FMT_USD)
        if not pd.isna(h.cost_basis):
            inp(ws, f"G{row}", h.cost_basis, FMT_USD)
        calc(ws, f"H{row}", f'=IF(ISNUMBER(G{row}),F{row}-G{row},"n/a")', FMT_USD)
        calc(ws, f"I{row}", f'=IF(N(G{row})>0,F{row}/G{row}-1,"n/a")', FMT_PCT)
        calc(ws, f"J{row}", f"=F{row}/TotalValue", FMT_PCT)
        for col, value in (("K", r.dcf_value), ("L", r.comps_value), ("M", r.target)):
            if not pd.isna(value):
                inp(ws, f"{col}{row}", value, FMT_PER_SHARE)
        calc(ws, f"N{row}", f'=IF(AND(N(M{row})>0,N(E{row})>0),M{row}/E{row}-1,"n/m")', FMT_PCT)
        # Helpers for the value-weighted upside: plain SUMs of these, so the math is auditable.
        calc(ws, f"O{row}", f"=IF(ISNUMBER(N{row}),F{row}*N{row},0)", FMT_USD)
        calc(ws, f"P{row}", f"=IF(ISNUMBER(N{row}),F{row},0)", FMT_USD)
        if r.workbook:
            c = ws[f"Q{row}"]
            c.value, c.hyperlink = "Open model", r.workbook.name     # relative link: same folder
            c.font = font("0563C1")
        label(ws, f"R{row}", r.reason, italic=True, color=GREY)

    # Totals
    t = last + 1
    label(ws, f"A{t}", "Total", bold=True)
    calc(ws, f"F{t}", f"=SUM(F{first}:F{last})", FMT_USD, bold=True)
    calc(ws, f"G{t}", f"=SUM(G{first}:G{last})", FMT_USD, bold=True)
    calc(ws, f"H{t}", f"=SUM(H{first}:H{last})", FMT_USD, bold=True)          # text "n/a" is ignored
    calc(ws, f"I{t}", f'=IF(G{t}>0,H{t}/G{t},"n/a")', FMT_PCT, bold=True)   # unknown-cost rows excluded from both
    calc(ws, f"J{t}", f"=SUM(J{first}:J{last})", FMT_PCT, bold=True)
    for col in "ABCDEFGHIJKLMNOPQR":
        ws[f"{col}{t}"].border = THIN_TOP
    wb.defined_names["TotalValue"] = DefinedName("TotalValue", attr_text=f"Portfolio!$F${t}")

    s = t + 2
    stats = [
        ("Value with a meaningful model target", f"=SUM(P{first}:P{last})",
         FMT_USD),
        ("Share of portfolio", f"=IF(TotalValue>0,C{s}/TotalValue,0)", FMT_PCT),
        ("Value-weighted model upside (those holdings)",
         f'=IF(C{s}>0,SUM(O{first}:O{last})/C{s},"n/m")', FMT_PCT),
    ]
    for i, (text, formula, fmt) in enumerate(stats):
        label(ws, f"A{s + i}", text, bold=i == 2)
        cell = calc(ws, f"C{s + i}", formula, fmt, bold=i == 2)
        if i == 2:
            cell.fill = KEY_FILL
    label(ws, f"A{s + 4}", "Blue = from the Fidelity export or each holding's model; black = formula. Gain % and totals "
          "exclude holdings whose cost basis Fidelity reports as \"--\". Upside compares the model target with the "
          "exported price; it is n/m when the target is $0 or below (a share can't be worth less than zero).", italic=True, color=GREY)

    widths(ws, A=18, B=40, C=12, D=10, E=11, F=13, G=13, H=12, I=9, J=8, K=12, L=13, M=14, N=14, O=15, P=15, Q=12, R=58)
    ws.freeze_panes = "B5"

    _history_sheet(wb.create_sheet("History"), history, totals)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _history_sheet(ws, history: pd.DataFrame, totals: pd.DataFrame):
    title(ws, "A1", "History – the portfolio across runs",
          "One snapshot per run date (re-running on the same day replaces that day's snapshot)")
    header_row(ws, 4, 1, ["Date", "Market value", "Cost basis (known)", "Gain $", "Gain %", "Value w/ model target",
                          "Weighted model upside"])
    fmts = [None, FMT_USD, FMT_USD, FMT_USD, FMT_PCT, FMT_USD, FMT_PCT]
    cols = ["date", "market_value", "cost_basis_known", "gain", "gain_pct", "modeled_value", "weighted_upside"]
    r = 5
    for _, row in totals.iterrows():
        for j, (col, fmt) in enumerate(zip(cols, fmts)):
            v = row[col]
            if not (isinstance(v, float) and pd.isna(v)):
                inp(ws, ws.cell(row=r, column=1 + j).coordinate, v, fmt)
        r += 1

    r += 2
    detail_cols = ["date", "symbol", "status", "quantity", "price", "market_value", "cost_basis", "gain",
                   "dcf_value", "comps_value", "target", "upside"]
    header_row(ws, r, 1, ["Date", "Symbol", "Status", "Quantity", "Price", "Market value", "Cost basis", "Gain $",
                          "DCF value", "Comps value", "Model target", "Model upside"])
    dfmts = [None, None, None, FMT_QTY, FMT_USD, FMT_USD, FMT_USD, FMT_USD,
             FMT_PER_SHARE, FMT_PER_SHARE, FMT_PER_SHARE, FMT_PCT]
    for _, row in history.sort_values(["date", "symbol"]).iterrows():
        r += 1
        for j, (col, fmt) in enumerate(zip(detail_cols, dfmts)):
            v = row[col]
            if not (isinstance(v, float) and pd.isna(v)):
                inp(ws, ws.cell(row=r, column=1 + j).coordinate, v, qty_format(v) if col == "quantity" else fmt)
    for c in ws["A"]:
        c.alignment = Alignment(horizontal="left")
    widths(ws, A=12, B=18, C=18, D=13, E=11, F=14, G=23, H=12, I=12, J=12, K=13, L=13)
    ws.freeze_panes = "A5"
