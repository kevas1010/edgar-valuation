"""Build the valuation workbook.

Only raw data (EDGAR values, prices, regression output) and assumptions are
written as values. Everything downstream is a live Excel formula, so changing
any blue input on the Assumptions sheet flows through to the Summary.

Key scalars are Excel *named ranges* (WACC, TerminalGrowth, SharePrice, ...)
so formulas read like the finance: =TerminalFCF*(1+g)/(WACC-g).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from edgar.client import FACTS_URL
from edgar.extract import DataNote, Financials
from model.assumptions import Assumptions, Scenario
from model.beta import BetaResult
from model.comps import MAX_MULTIPLE, PeerData
from model.prices import MARKET_INDEX
from excel.styles import (
    BLUE, FMT_FACTOR, FMT_PCT2, FMT_INT, FMT_MM, FMT_MULT, FMT_NUM, FMT_NUM4, FMT_PCT, FMT_PER_SHARE, FMT_SHARES_MM,
    GREY, KEY_FILL, THIN_TOP, calc, font, header_row, inp, label, link, section, title, widths,
)

SHEETS = ["Summary", "Historicals", "Assumptions", "DCF", "Sensitivity", "Beta Regression",
          "Comps", "Scenarios", "Data Notes"]
RAW_ROWS = ["Revenue", "Operating Income", "Net Income", "D&A", "Operating Cash Flow", "CapEx",
            "Pretax Income", "Income Tax", "Cash", "Debt"]
PROJECTION_YEARS = 5
YEAR_COLS = [get_column_letter(3 + k) for k in range(PROJECTION_YEARS)]     # C..G = Years 1-5 on DCF/Scenarios
FIRST, LAST = YEAR_COLS[0], YEAR_COLS[-1]
SENS_WACC_STEPS = [x / 1000 for x in range(-20, 21, 5)]                      # ±2% in 0.5% steps
SENS_GROWTH = [0.02, 0.025, 0.03, 0.035, 0.04]


@dataclass
class ModelData:
    fin: Financials
    assumptions: Assumptions
    scenarios: list[Scenario]
    beta: BetaResult
    price: float
    price_date: str
    as_of: date
    peers: list[PeerData] = field(default_factory=list)
    peer_tickers: list[str] = field(default_factory=list)   # as requested on the command line
    pulled_at: datetime = field(default_factory=datetime.now)


def _if_both(a: str, b: str, expr: str) -> str:
    """Formula that returns "" unless both referenced cells hold numbers (missing data stays blank)."""
    return f'=IF(COUNT({a},{b})=2,{expr},"")'


class ValuationWorkbook:
    def __init__(self, data: ModelData):
        self.d = data
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws = {name: self.wb.create_sheet(name) for name in SHEETS}
        self.hist_rows: dict[str, int] = {}
        self.hist_latest_col = ""

    # ------------------------------------------------------------------ helpers
    def define(self, name: str, sheet: str, ref: str) -> None:
        """Create a workbook-level named range pointing at a single cell."""
        col = "".join(ch for ch in ref if ch.isalpha())
        row = "".join(ch for ch in ref if ch.isdigit())
        self.wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${col}${row}")

    def hist(self, item: str, col: str | None = None) -> str:
        """Cross-sheet reference to a Historicals cell (latest fiscal year by default)."""
        return f"Historicals!{col or self.hist_latest_col}{self.hist_rows[item]}"

    @staticmethod
    def block(ws, start_row: int, spec: list, value_col: str = "B") -> dict[str, int]:
        """Write a vertical label/value block. `spec` rows are (key, label, make_formula(R), fmt, kind, bold),
        where R maps every key to its row number, so formulas can reference siblings by name."""
        R = {row[0]: start_row + i for i, row in enumerate(spec)}
        for key, text, make, fmt, kind, bold in spec:
            label(ws, f"A{R[key]}", text.strip(), bold=bold, italic=text.startswith("  "),
                  indent=1 if text.startswith("  ") else 0)
            kind(ws, f"{value_col}{R[key]}", make(R), fmt, bold=bold)
        return R

    def build(self, path: Path | str) -> Path:
        # Order matters: later sheets reference rows recorded by earlier ones.
        self._historicals()
        self._assumptions()
        self._beta()
        self._dcf()
        self._sensitivity()
        self._comps()
        self._scenarios()
        self._data_notes()
        self._summary()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        return path

    # ------------------------------------------------------------------ Historicals
    def _historicals(self):
        ws, fin = self.ws["Historicals"], self.d.fin
        years = list(fin.table.columns)
        cols = [get_column_letter(2 + i) for i in range(len(years))]
        src_col = get_column_letter(2 + len(years))
        self.hist_latest_col = cols[-1]

        title(ws, "A1", f"{fin.entity_name} – Historical Financials",
              f"$ millions · fiscal years ending ~{fin.year_ends[years[-1]]:%B %d} · "
              "source: SEC EDGAR XBRL companyfacts (10-K, FY)")
        header_row(ws, 4, 1, ["($ millions)"] + [f"FY{y}" for y in years] + ["Source XBRL tag (us-gaap)"])

        r = 5
        for item in RAW_ROWS:
            label(ws, f"A{r}", item)
            for c, y in zip(cols, years):
                v = fin.table.loc[item, y]
                if pd.notna(v):              # missing stays a truly empty cell (see Data Notes)
                    inp(ws, f"{c}{r}", float(v), FMT_MM)
            label(ws, f"{src_col}{r}", fin.sources.get(item, ""), italic=True, color=GREY)
            self.hist_rows[item] = r
            r += 1

        r += 1
        section(ws, r, 1, len(years) + 2, "Derived metrics (formulas)")
        r += 1
        H = self.hist_rows

        def derived(name: str, fmt: str, make, first_col: int = 0):
            nonlocal r
            label(ws, f"A{r}", name)
            for i, c in enumerate(cols):
                if i >= first_col:
                    calc(ws, f"{c}{r}", make(i, c), fmt)
            H[name] = r
            r += 1

        def ratio(num: str, den: str):
            return lambda i, c: _if_both(f"{c}{H[num]}", f"{c}{H[den]}", f'IFERROR({c}{H[num]}/{c}{H[den]},"")')

        derived("EBITDA", FMT_MM, lambda i, c: _if_both(f"{c}{H['Operating Income']}", f"{c}{H['D&A']}",
                                                         f"{c}{H['Operating Income']}+{c}{H['D&A']}"))
        derived("Free Cash Flow", FMT_MM, lambda i, c: _if_both(f"{c}{H['Operating Cash Flow']}", f"{c}{H['CapEx']}",
                                                                 f"{c}{H['Operating Cash Flow']}-{c}{H['CapEx']}"))
        derived("Revenue Growth %", FMT_PCT, lambda i, c: _if_both(
            f"{cols[i - 1]}{H['Revenue']}", f"{c}{H['Revenue']}",
            f'IFERROR({c}{H["Revenue"]}/{cols[i - 1]}{H["Revenue"]}-1,"")'), first_col=1)
        derived("EBITDA Margin %", FMT_PCT, ratio("EBITDA", "Revenue"))
        derived("FCF Margin %", FMT_PCT, ratio("Free Cash Flow", "Revenue"))
        derived("D&A % of Revenue", FMT_PCT, ratio("D&A", "Revenue"))
        derived("CapEx % of Revenue", FMT_PCT, ratio("CapEx", "Revenue"))
        derived("Effective Tax Rate %", FMT_PCT, ratio("Income Tax", "Pretax Income"))

        widths(ws, A=24, **{c: 12 for c in cols}, **{src_col: 70})
        ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ Assumptions
    def _assumptions(self):
        ws, a, fin = self.ws["Assumptions"], self.d.assumptions, self.d.fin
        notes_col = get_column_letter(2 + PROJECTION_YEARS)
        title(ws, "A1", "Assumptions",
              "Blue = input (edit these) · Black = formula · Green = link to another sheet")
        header_row(ws, 4, 1, ["Operating assumptions"] + [f"Year {k}" for k in range(1, PROJECTION_YEARS + 1)]
                   + ["Basis for default"])

        label(ws, "A5", "Revenue growth %")
        for k, g in enumerate(a.revenue_growth):
            inp(ws, f"{get_column_letter(2 + k)}5", g, FMT_PCT)
        label(ws, f"{notes_col}5", a.basis["revenue_growth"], italic=True, color=GREY)
        self.growth_row = 5

        def scalar(row: int, text: str, value, fmt, name: str, basis: str, is_link: bool = False):
            label(ws, f"A{row}", text)
            (link if is_link else inp)(ws, f"B{row}", value, fmt)
            label(ws, f"{notes_col}{row}", basis, italic=True, color=GREY)
            self.define(name, "Assumptions", f"B{row}")

        scalar(6, "EBITDA margin %", a.ebitda_margin, FMT_PCT, "EBITDAMargin", a.basis["ebitda_margin"])
        scalar(7, "D&A % of revenue", a.da_pct, FMT_PCT, "DAPct", a.basis["da_pct"])
        scalar(8, "CapEx % of revenue", a.capex_pct, FMT_PCT, "CapExPct", a.basis["capex_pct"])
        scalar(9, "Tax rate", a.tax_rate, FMT_PCT, "TaxRate", a.basis["tax_rate"])

        section(ws, 11, 1, 2 + PROJECTION_YEARS, "Cost of capital & terminal value")
        scalar(12, "Risk-free rate", a.risk_free, FMT_PCT2, "RiskFree", a.basis["risk_free"])
        scalar(13, "Equity risk premium", a.equity_risk_premium, FMT_PCT2, "EquityRiskPremium",
               a.basis["equity_risk_premium"])
        scalar(14, "Beta", a.beta, FMT_NUM, "EquityBeta", a.basis["beta"])
        scalar(15, "Pre-tax cost of debt", a.cost_of_debt, FMT_PCT, "CostOfDebt", a.basis["cost_of_debt"])
        scalar(16, "Debt materiality threshold", a.debt_materiality, FMT_PCT, "DebtThreshold",
               a.basis["debt_materiality"])
        scalar(17, "Terminal growth rate", a.terminal_growth, FMT_PCT, "TerminalGrowth", a.basis["terminal_growth"])
        scalar(18, "Terminal-year CapEx = D&A?", "Yes" if a.terminal_capex_equals_da else "No", None,
               "TerminalCapexEqDA", a.basis["terminal_capex_equals_da"])
        ws["B18"].alignment = Alignment(horizontal="right")
        yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
        ws.add_data_validation(yes_no)
        yes_no.add("B18")

        section(ws, 20, 1, 2 + PROJECTION_YEARS, "Market data & balance sheet")
        scalar(21, "Current share price", self.d.price, FMT_PER_SHARE, "SharePrice",
               f"yfinance close, {self.d.price_date}")
        scalar(22, "Shares outstanding (millions)", fin.shares_outstanding, FMT_SHARES_MM, "SharesOut",
               f"{fin.sources.get('Shares Outstanding', '')}, as of {fin.shares_as_of}")
        scalar(23, "Debt ($ millions)", f"={self.hist('Debt')}", FMT_MM, "BookDebt",
               f"Latest fiscal year (FY{fin.table.columns[-1]}) balance sheet, from Historicals", is_link=True)
        scalar(24, "Cash ($ millions)", f"={self.hist('Cash')}", FMT_MM, "Cash",
               f"Latest fiscal year (FY{fin.table.columns[-1]}) balance sheet, from Historicals", is_link=True)

        widths(ws, A=30, **{get_column_letter(2 + k): 11 for k in range(PROJECTION_YEARS)}, **{notes_col: 80})
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ DCF
    def _dcf(self):
        ws, fin = self.ws["DCF"], self.d.fin
        base_fy = int(fin.table.columns[-1])
        cols = ["B"] + YEAR_COLS                     # B = base year actuals, C..G = Years 1-5

        title(ws, "A1", "Discounted Cash Flow (DCF)",
              "$ millions except per share · unlevered free cash flow · end-of-year discounting")
        header_row(ws, 4, 1, ["Fiscal year", f"FY{base_fy}A"] + [f"FY{base_fy + k}E" for k in range(1, 6)])

        label(ws, "A5", "Period (t)", italic=True)
        calc(ws, "B5", 0, FMT_INT)
        for prev, c in zip(cols, YEAR_COLS):
            calc(ws, f"{c}5", f"={prev}5+1", FMT_INT)

        # Base-year actuals (green links) for context.
        link(ws, "B6", f"={self.hist('Revenue')}", FMT_MM)
        link(ws, "B7", f"={self.hist('Revenue Growth %')}", FMT_PCT)
        link(ws, "B8", f"={self.hist('EBITDA')}", FMT_MM)
        link(ws, "B9", f"={self.hist('EBITDA Margin %')}", FMT_PCT)
        link(ws, "B10", f"={self.hist('D&A')}", FMT_MM)
        link(ws, "B15", f"={self.hist('CapEx')}", FMT_MM)

        rows = [
            (6, "Revenue", FMT_MM, lambda p, c, k: f"={p}6*(1+Assumptions!{get_column_letter(1 + k)}${self.growth_row})"),
            (7, "  % growth", FMT_PCT, lambda p, c, k: f"=IFERROR({c}6/{p}6-1,0)"),
            (8, "EBITDA", FMT_MM, lambda p, c, k: f"={c}6*EBITDAMargin"),
            (9, "  % margin", FMT_PCT, lambda p, c, k: f"=IFERROR({c}8/{c}6,0)"),
            (10, "Less: D&A", FMT_MM, lambda p, c, k: f"={c}6*DAPct"),
            (11, "EBIT", FMT_MM, lambda p, c, k: f"={c}8-{c}10"),
            (12, "Less: taxes on EBIT (none on losses)", FMT_MM, lambda p, c, k: f"=MAX({c}11,0)*TaxRate"),
            (13, "NOPAT", FMT_MM, lambda p, c, k: f"={c}11-{c}12"),
            (14, "Plus: D&A", FMT_MM, lambda p, c, k: f"={c}10"),
            (15, "Less: CapEx", FMT_MM, lambda p, c, k: f"={c}6*CapExPct"),
            (16, "Unlevered free cash flow", FMT_MM, lambda p, c, k: f"={c}13+{c}14-{c}15"),
            (17, "  % of revenue", FMT_PCT, lambda p, c, k: f"=IFERROR({c}16/{c}6,0)"),
            (18, "Discount factor", FMT_FACTOR, lambda p, c, k: f"=1/(1+WACC)^{c}$5"),
            (19, "PV of unlevered FCF", FMT_MM, lambda p, c, k: f"={c}16*{c}18"),
        ]
        for row, text, fmt, make in rows:
            bold = row in (16, 19)
            label(ws, f"A{row}", text.strip(), bold=bold, italic=text.startswith("  "),
                  indent=1 if text.startswith("  ") else 0)
            for k, (p, c) in enumerate(zip(cols, YEAR_COLS), start=1):
                cell = calc(ws, f"{c}{row}", make(p, c, k), fmt, bold=bold)
                if row == 16:
                    cell.border = THIN_TOP

        # --- WACC (CAPM) ---
        section(ws, 21, 1, 7, "Cost of capital (WACC)")
        W = self.block(ws, 22, [
            ("rf", "Risk-free rate", lambda R: "=RiskFree", FMT_PCT2, link, False),
            ("beta", "Beta", lambda R: "=EquityBeta", FMT_NUM, link, False),
            ("erp", "Equity risk premium", lambda R: "=EquityRiskPremium", FMT_PCT2, link, False),
            ("ke", "Cost of equity (CAPM: rf + β × ERP)",
             lambda R: f"=B{R['rf']}+B{R['beta']}*B{R['erp']}", FMT_PCT2, calc, False),
            ("kd_pre", "Pre-tax cost of debt", lambda R: "=CostOfDebt", FMT_PCT, link, False),
            ("tax", "Tax rate", lambda R: "=TaxRate", FMT_PCT, link, False),
            ("kd", "After-tax cost of debt", lambda R: f"=B{R['kd_pre']}*(1-B{R['tax']})", FMT_PCT, calc, False),
            ("mcap", "Market value of equity (price × shares)", lambda R: "=SharePrice*SharesOut", FMT_MM, calc, False),
            ("debt", "Book value of debt", lambda R: "=BookDebt", FMT_MM, link, False),
            ("dshare", "Debt / total capital",
             lambda R: f"=IFERROR(B{R['debt']}/(B{R['mcap']}+B{R['debt']}),0)", FMT_PCT, calc, False),
            ("thr", "Materiality threshold", lambda R: "=DebtThreshold", FMT_PCT, link, False),
            ("wd", "Debt weight used (0 if immaterial)",
             lambda R: f"=IF(B{R['dshare']}>=B{R['thr']},B{R['dshare']},0)", FMT_PCT, calc, False),
            ("we", "Equity weight", lambda R: f"=1-B{R['wd']}", FMT_PCT, calc, False),
            ("wacc", "WACC", lambda R: f"=B{R['we']}*B{R['ke']}+B{R['wd']}*B{R['kd']}", FMT_PCT2, calc, True),
        ])
        ws[f"B{W['wacc']}"].fill = KEY_FILL
        self.define("WACC", "DCF", f"B{W['wacc']}")

        # --- Terminal value, EV, equity value per share ---
        top = W["wacc"] + 2
        section(ws, top, 1, 7, "Enterprise value → equity value per share")
        V = self.block(ws, top + 1, [
            ("sum_pv", "Sum of PV of FCF (Years 1–5)", lambda R: f"=SUM({FIRST}19:{LAST}19)", FMT_MM, calc, False),
            ("g", "Terminal growth rate", lambda R: "=TerminalGrowth", FMT_PCT, link, False),
            ("tfcf", "Terminal-year FCF (NOPAT + D&A − CapEx*)",
             lambda R: f'={LAST}13+{LAST}14-IF(TerminalCapexEqDA="Yes",{LAST}14,{LAST}15)', FMT_MM, calc, False),
            ("tv", "Terminal value: FCF × (1 + g) / (WACC − g)",
             lambda R: f"=B{R['tfcf']}*(1+B{R['g']})/(WACC-B{R['g']})", FMT_MM, calc, False),
            ("pv_tv", "PV of terminal value", lambda R: f"=B{R['tv']}*{LAST}18", FMT_MM, calc, False),
            ("ev", "Enterprise value", lambda R: f"=B{R['sum_pv']}+B{R['pv_tv']}", FMT_MM, calc, True),
            ("pct_tv", "  % of EV from terminal value", lambda R: f"=IFERROR(B{R['pv_tv']}/B{R['ev']},0)",
             FMT_PCT, calc, False),
            ("exit", "  Implied terminal EV / EBITDA", lambda R: f'=IF({LAST}8>0,B{R["tv"]}/{LAST}8,"n/m")',
             FMT_MULT, calc, False),
            ("debt", "Less: debt", lambda R: "=BookDebt", FMT_MM, link, False),
            ("cash", "Plus: cash", lambda R: "=Cash", FMT_MM, link, False),
            ("equity", "Equity value", lambda R: f"=B{R['ev']}-B{R['debt']}+B{R['cash']}", FMT_MM, calc, True),
            ("shares", "Shares outstanding (millions)", lambda R: "=SharesOut", FMT_SHARES_MM, link, False),
            ("vps", "DCF value per share", lambda R: f"=B{R['equity']}/B{R['shares']}", FMT_PER_SHARE, calc, True),
            ("check", "Check: WACC > terminal growth",
             lambda R: '=IF(WACC>TerminalGrowth,"OK","ERROR – WACC must exceed g")', None, calc, False),
        ])
        ws[f"B{V['vps']}"].fill = KEY_FILL
        label(ws, f"A{V['check'] + 1}", "* CapEx = D&A in the terminal year when the Assumptions switch is \"Yes\".",
              italic=True, color=GREY)
        for name, key in [("TerminalFCF", "tfcf"), ("EnterpriseValue", "ev"), ("PctEVFromTV", "pct_tv"),
                          ("EquityValue", "equity"), ("DCFValue", "vps"), ("ModelCheck", "check")]:
            self.define(name, "DCF", f"B{V[key]}")

        widths(ws, A=44, **{c: 13 for c in cols})
        ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ Sensitivity
    def _sensitivity(self):
        ws = self.ws["Sensitivity"]
        title(ws, "A1", "Sensitivity – DCF value per share",
              "Rows: WACC (base ± 2%) · Columns: terminal growth · every cell re-runs the DCF with live formulas")
        g_cols = [get_column_letter(3 + i) for i in range(len(SENS_GROWTH))]
        label(ws, "C4", "Terminal growth rate →", bold=True)
        label(ws, "B5", "WACC ↓", bold=True)
        ws["B5"].alignment = Alignment(horizontal="center")
        for c, g in zip(g_cols, SENS_GROWTH):
            inp(ws, f"{c}5", g, FMT_PCT, bold=True)

        ufcf, periods = f"DCF!${FIRST}$16:${LAST}$16", f"DCF!${FIRST}$5:${LAST}$5"
        first_row = 6
        for i, step in enumerate(SENS_WACC_STEPS):
            r = first_row + i
            calc(ws, f"B{r}", f"=WACC{step:+.3f}", FMT_PCT, bold=True)
            for c in g_cols:
                w, g = f"$B{r}", f"{c}$5"
                calc(ws, f"{c}{r}",
                     f'=IF({w}<={g},"n/m",(SUMPRODUCT({ufcf},1/(1+{w})^{periods})'
                     f"+TerminalFCF*(1+{g})/({w}-{g})/(1+{w})^DCF!${LAST}$5-BookDebt+Cash)/SharesOut)",
                     FMT_PER_SHARE)
        last_row = first_row + len(SENS_WACC_STEPS) - 1

        grid = f"{g_cols[0]}{first_row}:{g_cols[-1]}{last_row}"
        ws.conditional_formatting.add(grid, ColorScaleRule(
            start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
            mid_color="FFEB84", end_type="max", end_color="63BE7B"))

        # Outline the base case (WACC + 0, g = 3%).
        base = ws[f"{g_cols[SENS_GROWTH.index(0.03)]}{first_row + SENS_WACC_STEPS.index(0.0)}"]
        thick = Side(style="medium")
        base.border = Border(left=thick, right=thick, top=thick, bottom=thick)
        base.font = font(bold=True)

        note_row = last_row + 2
        label(ws, f"B{note_row}", "Current price")
        link(ws, f"C{note_row}", "=SharePrice", FMT_PER_SHARE)
        label(ws, f"B{note_row + 1}", "Outlined cell = base WACC at 3% terminal growth; it equals the DCF value "
              "whenever Assumptions terminal growth is 3%.", italic=True, color=GREY)
        label(ws, f"B{note_row + 2}", '"n/m" = WACC ≤ terminal growth, where the Gordon growth formula breaks down.',
              italic=True, color=GREY)
        widths(ws, A=2, B=14, **{c: 13 for c in g_cols})
        ws.freeze_panes = "A6"

    # ------------------------------------------------------------------ Beta Regression
    def _beta(self):
        ws, b, ticker = self.ws["Beta Regression"], self.d.beta, self.d.fin.ticker
        ret = b.returns
        n = len(ret)
        first, last = 5, 4 + n
        y, x = f"$J${first}:$J${last}", f"$K${first}:$K${last}"

        title(ws, "A1", "Beta Regression",
              f"OLS: weekly {ticker} return = α + β × weekly S&P 500 return · "
              f"{ret.index[0]:%b %d, %Y} – {ret.index[-1]:%b %d, %Y}")

        header_row(ws, 4, 1, ["Statistic", "statsmodels (Python)", "Excel check"])
        stats = [
            (5, "Beta (slope)", b.beta, f"=SLOPE({y},{x})", FMT_NUM4),
            (6, "Alpha (weekly intercept)", b.alpha, f"=INTERCEPT({y},{x})", FMT_NUM4),
            (7, "R-squared", b.r_squared, f"=RSQ({y},{x})", FMT_NUM4),
            (8, "Standard error of beta", b.std_err, f"=INDEX(LINEST({y},{x},TRUE,TRUE),2,1)", FMT_NUM4),
            (9, "Observations (weeks)", b.n_obs, f"=COUNT({y})", FMT_INT),
        ]
        for row, text, value, check, fmt in stats:
            label(ws, f"A{row}", text)
            inp(ws, f"B{row}", value, fmt)
            calc(ws, f"C{row}", check, fmt)
        derived = [
            (10, "95% CI – low (β − 1.96 × SE)", "=B5-1.96*B8"),
            (11, "95% CI – high (β + 1.96 × SE)", "=B5+1.96*B8"),
            (12, "Blume-adjusted beta (0.67β + 0.33)", "=0.67*B5+0.33"),
        ]
        for row, text, formula in derived:
            label(ws, f"A{row}", text)
            calc(ws, f"B{row}", formula, FMT_NUM4)
        label(ws, "A13", "Assumptions!Beta uses the raw beta (B5). Overwrite it there to use the adjusted beta instead.",
              italic=True, color=GREY)

        # Chart data: weekly returns, plus a two-point fit line driven by the stats above.
        header_row(ws, 4, 9, ["Week of", f"{ticker} return", "S&P 500 return"])
        for i, (dt, row) in enumerate(ret.iterrows()):
            r = first + i
            inp(ws, f"I{r}", dt.to_pydatetime(), "yyyy-mm-dd")
            inp(ws, f"J{r}", float(row["stock"]), "0.00%")
            inp(ws, f"K{r}", float(row["market"]), "0.00%")

        header_row(ws, 4, 13, ["Fit line: S&P 500", f"Fitted {ticker}"])
        calc(ws, "M5", f"=MIN({x})", "0.00%")
        calc(ws, "M6", f"=MAX({x})", "0.00%")
        for r in (5, 6):
            calc(ws, f"N{r}", f"=$B$6+$B$5*M{r}", "0.00%")

        chart = ScatterChart()
        chart.title = f"{ticker} vs S&P 500 weekly returns (β = {b.beta:.2f}, R² = {b.r_squared:.2f})"
        chart.style = 13
        chart.x_axis.title = "S&P 500 weekly return"
        chart.y_axis.title = f"{ticker} weekly return"
        chart.x_axis.number_format = chart.y_axis.number_format = "0%"
        chart.x_axis.delete = chart.y_axis.delete = False     # newer Excel hides axes otherwise
        # Axes cross at their minimums so labels and titles sit at the plot edges, not on the zero lines.
        chart.x_axis.crosses = chart.y_axis.crosses = "min"
        chart.legend.position = "r"
        # openpyxl 3.1 draws titles/legend on top of the plot unless overlay is explicitly off.
        for part in (chart.title, chart.x_axis.title, chart.y_axis.title, chart.legend):
            part.overlay = False

        points = Series(Reference(ws, min_col=10, min_row=first, max_row=last),
                        Reference(ws, min_col=11, min_row=first, max_row=last), title="Weekly returns")
        points.marker.symbol = "circle"
        points.marker.size = 5
        points.graphicalProperties.line.noFill = True
        fit = Series(Reference(ws, min_col=14, min_row=5, max_row=6),
                     Reference(ws, min_col=13, min_row=5, max_row=6), title="OLS fit")
        fit.marker.symbol = "none"
        fit.graphicalProperties.line.solidFill = "C00000"
        fit.graphicalProperties.line.width = 28575
        chart.series.extend([points, fit])
        chart.width, chart.height = 17, 10
        ws.add_chart(chart, "A15")

        widths(ws, A=36, B=20, C=14, D=3, I=12, J=13, K=15, L=3, M=17, N=14)
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ Comps
    def _comps(self):
        ws, peers, fin = self.ws["Comps"], self.d.peers, self.d.fin
        title(ws, "A1", "Comparable Companies",
              "$ millions except per share · latest fiscal year from SEC EDGAR · prices from yfinance · "
              '"n/m" = negative/missing metric or above the outlier cap, excluded from median/mean')
        label(ws, "J3", "Outlier cap – multiples above this are n/m:", italic=True)
        inp(ws, "O3", MAX_MULTIPLE, FMT_MULT)
        self.define("MaxMultiple", "Comps", "O3")
        header_row(ws, 4, 1, ["Ticker", "Company", "FY", "Price", "Shares (mm)", "Market cap", "Debt", "Cash",
                              "Enterprise value", "Revenue", "EBITDA", "Net income",
                              "EV / Revenue", "EV / EBITDA", "P / E"])
        first = 5
        for i, p in enumerate(peers):
            r = first + i
            label(ws, f"A{r}", p.ticker, bold=True)
            label(ws, f"B{r}", p.name)
            label(ws, f"C{r}", f"FY{p.fiscal_year}")
            inp(ws, f"D{r}", p.price, FMT_PER_SHARE)
            inp(ws, f"E{r}", p.shares, FMT_SHARES_MM)
            calc(ws, f"F{r}", f"=D{r}*E{r}", FMT_MM)
            inp(ws, f"G{r}", p.debt, FMT_MM)
            inp(ws, f"H{r}", p.cash, FMT_MM)
            calc(ws, f"I{r}", f"=F{r}+G{r}-H{r}", FMT_MM)
            for col, value in (("J", p.revenue), ("K", p.ebitda), ("L", p.net_income)):
                if pd.notna(value):
                    inp(ws, f"{col}{r}", value, FMT_MM)
            self._multiples(ws, r)
        last = first + len(peers) - 1

        r = last + 1
        for text, fn in (("Median", "MEDIAN"), ("Mean", "AVERAGE")):
            label(ws, f"A{r}", text, bold=True)
            for col in "MNO":
                cell = calc(ws, f"{col}{r}", f'=IFERROR({fn}({col}{first}:{col}{last}),"n/m")' if peers else '="n/m"',
                            FMT_MULT, bold=True)
                cell.border = THIN_TOP if text == "Median" else Border()
            r += 1
        median_row = last + 1

        # The target at today's price, for reference (not part of the median).
        r += 1
        label(ws, f"A{r}", fin.ticker, bold=True, color=BLUE)
        label(ws, f"B{r}", f"{fin.entity_name} – target", italic=True)
        label(ws, f"C{r}", f"FY{fin.table.columns[-1]}")
        link(ws, f"D{r}", "=SharePrice", FMT_PER_SHARE)
        link(ws, f"E{r}", "=SharesOut", FMT_SHARES_MM)
        calc(ws, f"F{r}", f"=D{r}*E{r}", FMT_MM)
        link(ws, f"G{r}", "=BookDebt", FMT_MM)
        link(ws, f"H{r}", "=Cash", FMT_MM)
        calc(ws, f"I{r}", f"=F{r}+G{r}-H{r}", FMT_MM)
        link(ws, f"J{r}", f"={self.hist('Revenue')}", FMT_MM)
        link(ws, f"K{r}", f"={self.hist('EBITDA')}", FMT_MM)
        link(ws, f"L{r}", f"={self.hist('Net Income')}", FMT_MM)
        self._multiples(ws, r)

        # Apply peer medians to the target's metrics.
        r += 2
        header_row(ws, r, 1, ["Target", "Method", "Metric", "Median", "Implied EV", "Implied equity", "Per share"])
        methods = [
            ("EV / Revenue", "Revenue", "M", True),
            ("EV / EBITDA", "EBITDA", "N", True),
            ("P / E", "Net Income", "O", False),
        ]
        per_share = []
        for name, metric, mcol, is_ev in methods:
            r += 1
            label(ws, f"A{r}", fin.ticker)
            label(ws, f"B{r}", name)
            link(ws, f"C{r}", f"={self.hist(metric)}", FMT_MM)
            link(ws, f"D{r}", f"={mcol}{median_row}", FMT_MULT)
            ok = f"AND(ISNUMBER(C{r}),ISNUMBER(D{r}),N(C{r})>0)"
            if is_ev:
                calc(ws, f"E{r}", f'=IF({ok},C{r}*D{r},"n/m")', FMT_MM)
                calc(ws, f"F{r}", f'=IF(ISNUMBER(E{r}),E{r}-BookDebt+Cash,"n/m")', FMT_MM)
            else:
                calc(ws, f"F{r}", f'=IF({ok},C{r}*D{r},"n/m")', FMT_MM)
            calc(ws, f"G{r}", f'=IF(ISNUMBER(F{r}),F{r}/SharesOut,"n/m")', FMT_PER_SHARE)
            per_share.append(f"G{r}")

        r += 1
        label(ws, f"A{r}", "Comps-implied value per share (average of meaningful methods)", bold=True)
        cell = calc(ws, f"G{r}", f'=IFERROR(AVERAGE({per_share[0]}:{per_share[-1]}),"n/m")', FMT_PER_SHARE, bold=True)
        cell.fill = KEY_FILL
        cell.border = THIN_TOP
        self.define("CompsValue", "Comps", f"G{r}")
        label(ws, f"A{r + 2}", "Methods are n/m when the target's own metric is negative (e.g. a loss-making year). "
              "Peer metrics are the latest fiscal year, not LTM, and fiscal year-ends differ across peers.",
              italic=True, color=GREY)
        if not peers:
            label(ws, f"A{first}", "No peers provided (use --peers).", italic=True, color=GREY)

        widths(ws, A=10, B=38, C=11, D=11, E=12, F=13, G=12, H=12, I=15, J=12, K=12, L=12, M=13, N=13, O=10)
        ws.freeze_panes = "C5"

    @staticmethod
    def _multiples(ws, r: int):
        """n/m when the metric is missing / non-positive, or the multiple exceeds the outlier cap.
        Nested IFs because Excel's AND() does not short-circuit (I/0 would still error)."""
        for col, num, den in (("M", "I", "J"), ("N", "I", "K"), ("O", "F", "L")):
            m = f"{num}{r}/{den}{r}"
            calc(ws, f"{col}{r}", f'=IF(N({den}{r})>0,IF({m}<=MaxMultiple,{m},"n/m"),"n/m")', FMT_MULT)

    # ------------------------------------------------------------------ Scenarios
    def _scenarios(self):
        ws, scen = self.ws["Scenarios"], self.d.scenarios
        base_fy = int(self.d.fin.table.columns[-1])
        title(ws, "A1", "Scenarios – probability-weighted price target",
              "Base links to the Assumptions sheet; Bear / Bull are their own blue inputs. "
              "Each case re-runs the full DCF below.")
        s_cols = ["B", "C", "D"]                              # Bear, Base, Bull

        header_row(ws, 4, 1, ["Scenario inputs"] + [s.name for s in scen] + ["Notes"])
        label(ws, "A5", "Probability", bold=True)
        for c, s in zip(s_cols, scen):
            inp(ws, f"{c}5", s.probability, FMT_PCT, bold=True)
        for k in range(PROJECTION_YEARS):
            r = 6 + k
            label(ws, f"A{r}", f"Revenue growth – Year {k + 1}")
            for c, s in zip(s_cols, scen):
                if s.name == "Base":
                    link(ws, f"{c}{r}", f"=Assumptions!{get_column_letter(2 + k)}{self.growth_row}", FMT_PCT)
                else:
                    inp(ws, f"{c}{r}", s.revenue_growth[k], FMT_PCT)
        label(ws, "A11", "EBITDA margin")
        for c, s in zip(s_cols, scen):
            if s.name == "Base":
                link(ws, f"{c}11", "=EBITDAMargin", FMT_PCT)
            else:
                inp(ws, f"{c}11", s.ebitda_margin, FMT_PCT)
        label(ws, "A12", "Probabilities sum to", italic=True)
        calc(ws, "B12", "=SUM(B5:D5)", FMT_PCT)
        calc(ws, "C12", '=IF(ROUND(B12,6)=1,"OK","Probabilities should sum to 100%")')
        label(ws, "E6", "Defaults: Bear / Bull = base ∓ 5 pts of growth each year", italic=True, color=GREY)
        label(ws, "E11", "Defaults: Bear / Bull = base ∓ 5 pts of margin", italic=True, color=GREY)
        label(ws, "E5", "D&A %, CapEx %, tax, WACC and terminal growth are shared with the base case",
              italic=True, color=GREY)

        # One full DCF per scenario.
        vps_cells = []
        top = 22
        for c, s in zip(s_cols, scen):
            vps_cells.append(self._scenario_block(ws, top, s.name, c, base_fy))
            top += 17

        header_row(ws, 14, 1, ["Scenario outputs"] + [s.name for s in scen])
        label(ws, "A15", "Value per share", bold=True)
        label(ws, "A16", "Upside / (downside) vs current price")
        label(ws, "A17", "Probability × value")
        for c, vps in zip(s_cols, vps_cells):
            calc(ws, f"{c}15", f"={vps}", FMT_PER_SHARE, bold=True)
            calc(ws, f"{c}16", f"=IFERROR({c}15/SharePrice-1,\"\")", FMT_PCT)
            calc(ws, f"{c}17", f"={c}5*{c}15", FMT_PER_SHARE)
        label(ws, "A19", "Probability-weighted price target", bold=True)
        cell = calc(ws, "B19", "=SUMPRODUCT(B5:D5,B15:D15)/SUM(B5:D5)", FMT_PER_SHARE, bold=True)
        cell.fill = KEY_FILL
        self.define("ScenarioTarget", "Scenarios", "B19")
        label(ws, "A20", "Check: base case = DCF sheet", italic=True)
        calc(ws, "B20", '=IF(ABS(C15-DCFValue)<0.005,"OK","Base differs from DCF")')

        widths(ws, A=40, B=13, C=13, D=13, E=13, F=13, G=13)
        ws.freeze_panes = "B5"

    def _scenario_block(self, ws, top: int, name: str, s_col: str, base_fy: int) -> str:
        """A compact DCF driven by one scenario's growth/margin column. Returns its value-per-share cell."""
        header_row(ws, top, 1, [f"{name} case ($ millions)", f"FY{base_fy}A"]
                   + [f"FY{base_fy + k}E" for k in range(1, 6)])
        R = {key: top + 1 + i for i, key in enumerate(
            ["rev", "ebitda", "da", "ebit", "tax", "nopat", "capex", "ufcf", "pv"])}
        labels = {"rev": "Revenue", "ebitda": "EBITDA", "da": "D&A", "ebit": "EBIT", "tax": "Taxes",
                  "nopat": "NOPAT", "capex": "CapEx", "ufcf": "Unlevered FCF", "pv": "PV of FCF"}
        for key, row in R.items():
            label(ws, f"A{row}", labels[key], bold=key == "ufcf")
        link(ws, f"B{R['rev']}", "=DCF!B6", FMT_MM)

        cols = ["B"] + YEAR_COLS
        for k, (p, c) in enumerate(zip(cols, YEAR_COLS), start=1):
            f = {
                "rev": f"={p}{R['rev']}*(1+${s_col}${5 + k})",
                "ebitda": f"={c}{R['rev']}*${s_col}$11",
                "da": f"={c}{R['rev']}*DAPct",
                "ebit": f"={c}{R['ebitda']}-{c}{R['da']}",
                "tax": f"=MAX({c}{R['ebit']},0)*TaxRate",
                "nopat": f"={c}{R['ebit']}-{c}{R['tax']}",
                "capex": f"={c}{R['rev']}*CapExPct",
                "ufcf": f"={c}{R['nopat']}+{c}{R['da']}-{c}{R['capex']}",
                "pv": f"={c}{R['ufcf']}/(1+WACC)^DCF!{c}$5",
            }
            for key, formula in f.items():
                calc(ws, f"{c}{R[key]}", formula, FMT_MM, bold=key == "ufcf")

        V = self.block(ws, R["pv"] + 1, [
            ("tfcf", "Terminal-year FCF",
             lambda X: f'={LAST}{R["nopat"]}+{LAST}{R["da"]}-IF(TerminalCapexEqDA="Yes",{LAST}{R["da"]},{LAST}{R["capex"]})',
             FMT_MM, calc, False),
            ("tv", "Terminal value", lambda X: f"=B{X['tfcf']}*(1+TerminalGrowth)/(WACC-TerminalGrowth)",
             FMT_MM, calc, False),
            ("ev", "Enterprise value",
             lambda X: f"=SUM({FIRST}{R['pv']}:{LAST}{R['pv']})+B{X['tv']}/(1+WACC)^DCF!{LAST}$5", FMT_MM, calc, False),
            ("equity", "Equity value", lambda X: f"=B{X['ev']}-BookDebt+Cash", FMT_MM, calc, False),
            ("vps", f"{name} value per share", lambda X: f"=B{X['equity']}/SharesOut", FMT_PER_SHARE, calc, True),
        ])
        return f"B{V['vps']}"

    # ------------------------------------------------------------------ Data Notes
    def _data_notes(self):
        ws, d, fin = self.ws["Data Notes"], self.d, self.d.fin
        title(ws, "A1", "Data Notes", "Where every number came from, and everything the extractor had to work around")
        header_row(ws, 4, 1, ["Run information", "", ""])
        info = [
            ("Data pulled", d.pulled_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Target", f"{fin.entity_name} ({fin.ticker})"),
            ("Financials source", "SEC EDGAR XBRL companyfacts API (10-K, fp = FY; restatements → latest filing)"),
            ("Companyfacts URL", FACTS_URL.format(cik=fin.cik)),
            ("Prices source", f"yfinance (current price as of {d.price_date}; 2y weekly adjusted closes vs {MARKET_INDEX})"),
            ("Peers requested", ", ".join(d.peer_tickers) or "none"),
            ("Peers used", ", ".join(p.ticker for p in d.peers) or "none"),
        ]
        r = 5
        for k, v in info:
            label(ws, f"A{r}", k, bold=True)
            label(ws, f"B{r}", v)
            r += 1

        r += 1
        header_row(ws, r, 1, ["Category", "Item", "Note"])
        notes = fin.notes or [DataNote("Info", "–", "No missing tags, fallbacks or assumptions to report.")]
        for n in notes:
            r += 1
            label(ws, f"A{r}", n.category)
            label(ws, f"B{r}", n.item)
            label(ws, f"C{r}", n.message)
            ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        widths(ws, A=18, B=26, C=120)
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ Summary
    def _summary(self):
        ws, fin = self.ws["Summary"], self.d.fin
        ws.sheet_view.showGridLines = False
        title(ws, "B2", f"{fin.entity_name} ({fin.ticker})",
              f"Equity valuation · prepared {self.d.as_of:%B %d, %Y} · $ millions except per share")

        header_row(ws, 5, 2, ["Valuation summary", "Per share", "vs. current"])
        rows = [
            (6, "Current share price", "=SharePrice", link),
            (7, "DCF value per share (base case)", "=DCFValue", link),
            (8, "Comps-implied value per share", "=CompsValue", link),
            (9, "Scenario-weighted price target", "=ScenarioTarget", link),
        ]
        for row, text, formula, kind in rows:
            bold = row == 9
            label(ws, f"B{row}", text, bold=bold)
            kind(ws, f"C{row}", formula, FMT_PER_SHARE, bold=bold)
            if row > 6:
                calc(ws, f"D{row}", f'=IF(ISNUMBER(C{row}),C{row}/$C$6-1,"n/m")', FMT_PCT, bold=bold)
        for c in ("B9", "C9", "D9"):
            ws[c].fill = KEY_FILL
            ws[c].border = THIN_TOP

        header_row(ws, 11, 2, ["Scenarios", "Per share", "Probability"])
        for i, (name, col) in enumerate((("Bear", "B"), ("Base", "C"), ("Bull", "D"))):
            r = 12 + i
            label(ws, f"B{r}", name)
            link(ws, f"C{r}", f"=Scenarios!{col}15", FMT_PER_SHARE)
            link(ws, f"D{r}", f"=Scenarios!{col}5", FMT_PCT)

        header_row(ws, 16, 2, ["Key outputs", "Value", ""])
        outputs = [
            ("Market capitalization", "=SharePrice*SharesOut", FMT_MM, calc),
            ("Enterprise value (DCF)", "=EnterpriseValue", FMT_MM, link),
            ("Equity value (DCF)", "=EquityValue", FMT_MM, link),
            ("WACC", "=WACC", FMT_PCT, link),
            ("Terminal growth rate", "=TerminalGrowth", FMT_PCT, link),
            ("Beta", "=EquityBeta", FMT_NUM, link),
            ("% of EV from terminal value", "=PctEVFromTV", FMT_PCT, link),
            ("Terminal-year CapEx = D&A?", "=TerminalCapexEqDA", None, link),
            ("Model check", "=ModelCheck", None, link),
        ]
        for i, (text, formula, fmt, kind) in enumerate(outputs):
            r = 17 + i
            label(ws, f"B{r}", text)
            kind(ws, f"C{r}", formula, fmt)
            ws[f"C{r}"].alignment = Alignment(horizontal="right")

        label(ws, f"B{18 + len(outputs)}",
              "Change any blue cell on the Assumptions or Scenarios sheets and every number here updates.",
              italic=True, color=BLUE)
        widths(ws, A=2, B=36, C=16, D=14)
        ws.freeze_panes = "A6"


def write_workbook(data: ModelData, path: Path | str) -> Path:
    return ValuationWorkbook(data).build(path)
