"""Banker formatting conventions shared by every sheet.

Colors:  blue = hardcoded input,  black = formula,  green = link to another sheet.
Money is stored in full dollars and *displayed* in millions via the ",," format
suffix, so per-share math never has to undo a /1e6.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
BLUE, BLACK, GREEN, WHITE, GREY = "0000FF", "000000", "008000", "FFFFFF", "808080"

FMT_MM = '$#,##0.0,,;($#,##0.0,,);"–"'        # $ millions, 1 decimal
FMT_PCT = "0.0%;(0.0%);0.0%"
FMT_PCT2 = "0.00%;(0.00%);0.00%"            # rates where 1 decimal would misstate the input
FMT_PER_SHARE = '$#,##0.00;($#,##0.00);"–"'
FMT_SHARES_MM = '#,##0.0,,'                   # share count in millions
FMT_MULT = '0.0"x";(0.0"x");"–"'
FMT_NUM = "0.00"
FMT_NUM4 = "0.0000"
FMT_INT = "0"
FMT_FACTOR = "0.000"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_TOP = Border(top=Side(style="thin"))


def font(color: str = BLACK, bold: bool = False, italic: bool = False, size: int = 10) -> Font:
    return Font(name=FONT_NAME, color=color, bold=bold, italic=italic, size=size)


def _put(ws: Worksheet, ref: str, value, color: str, fmt: str | None, bold: bool):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold)
    if fmt:
        c.number_format = fmt
    return c


def inp(ws, ref, value, fmt=None, bold=False):
    """Hardcoded input (blue)."""
    return _put(ws, ref, value, BLUE, fmt, bold)


def calc(ws, ref, formula, fmt=None, bold=False):
    """Formula computed on this sheet (black). Right-aligned so text results like "n/m" line up with numbers."""
    c = _put(ws, ref, formula, BLACK, fmt, bold)
    c.alignment = Alignment(horizontal="right")
    return c


def link(ws, ref, formula, fmt=None, bold=False):
    """Pure reference to another sheet or named input (green)."""
    c = _put(ws, ref, formula, GREEN, fmt, bold)
    c.alignment = Alignment(horizontal="right")
    return c


def label(ws, ref, text, bold=False, italic=False, indent=0, color=BLACK):
    c = ws[ref]
    c.value = text
    c.font = font(color, bold, italic)
    if indent:
        c.alignment = Alignment(indent=indent)
    return c


def title(ws, ref, text, subtitle: str | None = None):
    ws[ref].value = text
    ws[ref].font = font(bold=True, size=14)
    if subtitle:
        below = ws.cell(row=ws[ref].row + 1, column=ws[ref].column)
        below.value = subtitle
        below.font = font(GREY, italic=True)


def header_row(ws, row: int, first_col: int, values: list[str]):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=first_col + i, value=v)
        c.font = font(WHITE, bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if i else "left")


def section(ws, row: int, first_col: int, last_col: int, text: str):
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    c = ws.cell(row=row, column=first_col, value=text)
    c.font = font(bold=True)


def widths(ws, **cols: float):
    for col, w in cols.items():
        ws.column_dimensions[col].width = w
