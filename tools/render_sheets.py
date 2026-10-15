"""Render workbook sheets to PNG screenshots (for the README). Windows + Excel required.

Excel exports the workbook to PDF (so the images show exactly what Excel renders,
including conditional formatting and charts); PyMuPDF rasterizes the pages and
trims the white margins.

    python tools/render_sheets.py output/AEHR_valuation_2026-10-14.xlsx Summary Sensitivity
"""
from __future__ import annotations

import argparse
import subprocess
import tempfile
from pathlib import Path

import numpy as np
import pymupdf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "docs" / "images"


def trim_whitespace(pix: pymupdf.Pixmap, pad: int = 16) -> pymupdf.IRect:
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    ink = np.where((img[:, :, :3] < 245).any(axis=2))
    top, bottom = ink[0].min(), ink[0].max()
    left, right = ink[1].min(), ink[1].max()
    return pymupdf.IRect(max(left - pad, 0), max(top - pad, 0),
                         min(right + pad, pix.width), min(bottom + pad, pix.height))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("workbook", type=Path)
    p.add_argument("sheets", nargs="+", help="Sheet names to render, e.g. Summary Sensitivity")
    p.add_argument("--dpi", type=int, default=150)
    args = p.parse_args()

    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / "workbook.pdf"
        subprocess.run(["powershell", "-NoProfile", "-File", str(ROOT / "tools" / "export_pdf.ps1"),
                        "-Path", str(args.workbook.resolve()), "-Out", str(pdf)], check=True)
        doc = pymupdf.open(pdf)
        order = load_workbook(args.workbook, read_only=True).sheetnames   # one PDF page per sheet, in order
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        for sheet in args.sheets:
            page = doc[order.index(sheet)]
            pix = page.get_pixmap(dpi=args.dpi)
            clip = trim_whitespace(pix)
            scale = 72 / args.dpi
            rect = pymupdf.Rect(clip.x0 * scale, clip.y0 * scale, clip.x1 * scale, clip.y1 * scale)
            out = OUT_DIR / f"{sheet.lower().replace(' ', '_')}.png"
            page.get_pixmap(dpi=args.dpi, clip=rect).save(out)
            print(f"Wrote {out.relative_to(ROOT)}")
        doc.close()


if __name__ == "__main__":
    main()
