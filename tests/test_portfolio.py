"""Portfolio mode: Fidelity CSV parsing, per-holding valuation, the portfolio workbook, history.
Uses a FAKE positions file (tests/fixtures/fidelity_positions_sample.csv) - never real holdings."""
from __future__ import annotations

import math
from datetime import date

import pytest
from openpyxl import load_workbook

from excel.portfolio_writer import write_portfolio_workbook
from model.pipeline import build_model, python_values
from portfolio.fidelity import CASH, OPTION, STOCK, parse_money, parse_positions
from portfolio.run import MODELED, NOT_MODELED, append_history, history_totals, value_portfolio
from tests.conftest import FIXTURE_DATE, FIXTURES
from tests.test_dcf import Workbook

SAMPLE = FIXTURES / "fidelity_positions_sample.csv"


@pytest.fixture(scope="module")
def holdings():
    return {h.symbol: h for h in parse_positions(SAMPLE)}


@pytest.fixture(scope="module")
def run(holdings, client, price_loader, tmp_path_factory):
    out = tmp_path_factory.mktemp("portfolio")
    results = value_portfolio(list(holdings.values()), client, out, as_of=FIXTURE_DATE,
                              price_loader=price_loader, log=lambda *_: None)
    history = append_history(results, out / "history.csv", FIXTURE_DATE)
    totals = history_totals(history)
    path = write_portfolio_workbook(results, history, totals, out / "Portfolio_test.xlsx", FIXTURE_DATE)
    return {r.holding.symbol: r for r in results}, totals, path, out


# ---------------------------------------------------------------- parsing

def test_parse_money_handles_fidelity_formats():
    assert parse_money("$1,234.56") == 1234.56
    assert parse_money("-$12.30") == -12.30
    assert parse_money("+139.03%") == 139.03
    assert math.isnan(parse_money("--")) and math.isnan(parse_money(""))


def test_parser_combines_accounts_and_classifies(holdings):
    assert set(holdings) == {"SPAXX", "FDRXX", "AEHR", "TER", "FXAIX", "-AEHR261120C80", "COHU"}  # no Pending Activity
    aehr = holdings["AEHR"]                                  # 40 shares + 10 shares in two accounts
    assert (aehr.quantity, aehr.current_value, aehr.cost_basis) == (50, 2868.40 + 717.10, 1200 + 800)
    assert aehr.accounts == ["FAKE INDIVIDUAL", "FAKE ROTH IRA"]
    assert holdings["SPAXX"].kind == CASH and holdings["SPAXX"].current_value == 1250.0
    assert holdings["-AEHR261120C80"].kind == OPTION
    assert holdings["TER"].kind == STOCK
    assert math.isnan(holdings["COHU"].cost_basis)           # "--" cost basis -> unknown, not zero


def test_account_numbers_are_never_kept(holdings):
    assert "X0000000" not in repr(list(holdings.values()))


# ---------------------------------------------------------------- valuation

def test_stocks_are_modeled_and_everything_else_explained(run):
    results = run[0]
    assert {s for s, r in results.items() if r.status == MODELED} == {"AEHR", "TER", "COHU"}
    assert "mutual fund or ETF" in results["FXAIX"].reason
    assert results["SPAXX"].reason == "Cash / money market"
    assert results["-AEHR261120C80"].reason == "Option contract"
    assert all(results[s].workbook.exists() for s in ("AEHR", "TER", "COHU"))


def test_holding_target_matches_single_ticker_model(run, client, price_loader):
    """Portfolio mode must give the same answer as `valuate.py AEHR` (no peers) on the same data."""
    aehr = run[0]["AEHR"]
    expected = python_values(build_model("AEHR", [], client, price_loader, as_of=FIXTURE_DATE))
    assert aehr.target == pytest.approx(expected["Target"], rel=1e-12)
    assert aehr.upside == pytest.approx(expected["Target"] / 71.71 - 1, rel=1e-12)


# ---------------------------------------------------------------- workbook

def test_portfolio_workbook_formulas(run):
    results, totals, path, _ = run
    wb = Workbook(path)
    sol = wb.calculate()
    errors = {k: v for k, v in sol.items() if str(v).startswith("#")}
    assert not errors, f"Formula errors: {list(errors.items())[:10]}"

    total_value = sum(r.holding.current_value for r in results.values())
    assert wb.get(sol, "TotalValue") == pytest.approx(total_value)
    xl = load_workbook(path)["Portfolio"]
    rows = {xl[f"A{r}"].value: r for r in range(5, 5 + len(results))}
    total_row = 5 + len(results)
    known = [r.holding for r in results.values() if not math.isnan(r.holding.cost_basis)]
    assert wb.get(sol, f"Portfolio!H{total_row}") == pytest.approx(sum(h.gain for h in known))
    assert wb.get(sol, f"Portfolio!H{rows['COHU']}") == "n/a"
    # The workbook's value-weighted upside agrees with the History computation.
    assert wb.get(sol, f"Portfolio!C{total_row + 4}") == pytest.approx(totals.iloc[-1]["weighted_upside"], rel=1e-9)


def test_workbook_links_models_and_hides_account_numbers(run):
    path = run[2]
    ws = load_workbook(path)["Portfolio"]
    links = {ws[f"A{r}"].value: ws[f"Q{r}"].hyperlink for r in range(5, 12)}
    assert links["AEHR"].target.startswith("AEHR_valuation_") and links["SPAXX"] is None
    every_cell = " ".join(str(c.value) for sheet in load_workbook(path) for row in sheet.iter_rows() for c in row)
    assert "X0000000" not in every_cell


# ---------------------------------------------------------------- history

def test_history_replaces_same_day_and_appends_new_days(run):
    results, _, _, out = run
    hist_path = out / "history.csv"
    same_day = append_history(list(results.values()), hist_path, FIXTURE_DATE)
    assert len(same_day) == len(results)                         # re-run replaced, didn't duplicate
    next_day = append_history(list(results.values()), hist_path, date(2026, 10, 15))
    assert history_totals(next_day)["date"].tolist() == ["2026-10-14", "2026-10-15"]


def test_header_capitalization_does_not_matter(tmp_path):
    """Older Fidelity exports use Title Case headers ("Current Value"); newer ones use sentence case."""
    old_style = SAMPLE.read_text(encoding="utf-8-sig").replace("Current value", "Current Value") \
        .replace("Cost basis total", "Cost Basis Total").replace("Last price,", "Last Price,")
    path = tmp_path / "old_style.csv"
    path.write_text(old_style, encoding="utf-8")
    assert {h.symbol: h.current_value for h in parse_positions(path)} == \
           {h.symbol: h.current_value for h in parse_positions(SAMPLE)}


def test_negative_model_value_is_not_a_price_target(holdings):
    """A negative model value per share means 'equity worth nothing on these assumptions', not -150% upside.
    It keeps its DCF value but is excluded from the value-weighted upside (numerator and denominator)."""
    from portfolio.run import HoldingResult
    aehr, ter = holdings["AEHR"], holdings["TER"]
    good = HoldingResult(aehr, MODELED, target=aehr.last_price * 1.10)       # +10%
    bad = HoldingResult(ter, MODELED, target=-20.0)
    assert good.upside == pytest.approx(0.10) and math.isnan(bad.upside)
    totals = history_totals(append_history([good, bad], FIXTURES.parent / ".tmp_history.csv", FIXTURE_DATE))
    (FIXTURES.parent / ".tmp_history.csv").unlink()
    assert totals.iloc[0]["weighted_upside"] == pytest.approx(0.10)
    assert totals.iloc[0]["modeled_value"] == pytest.approx(aehr.current_value)


def test_company_without_revenue_is_not_modeled(holdings, client, price_loader, tmp_path, monkeypatch):
    """Holding partnerships and trusts (e.g. AB, ETHW) report net income but no revenue."""
    from portfolio import run as run_module
    real_build = run_module.build_model

    def no_revenue(*args, **kwargs):
        data = real_build(*args, **kwargs)
        data.fin.table.loc["Revenue"] = float("nan")
        return data
    monkeypatch.setattr(run_module, "build_model", no_revenue)
    r = run_module.value_holding(holdings["TER"], client, price_loader, FIXTURE_DATE, tmp_path)
    assert r.status == NOT_MODELED and "No revenue" in r.reason
