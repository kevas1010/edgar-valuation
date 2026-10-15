"""XBRL extraction: annual-only filtering, restatement dedupe, per-year tag fallbacks, missing data."""
from __future__ import annotations

import copy
import math

import pytest
from openpyxl import load_workbook

from edgar.extract import DataNote, annual_facts, extract_financials
from excel.writer import write_workbook


def fact(val, end, filed, form="10-K", fp="FY", start=None, fy=None):
    f = {"val": val, "end": end, "filed": filed, "form": form, "fp": fp, "fy": fy or int(end[:4]), "accn": filed}
    if start:
        f["start"] = start
    return f


def facts_with(tags: dict[str, list[dict]], shares: float = 1_000_000) -> dict:
    """Minimal companyfacts JSON: {tag: [facts]} under us-gaap, USD."""
    return {
        "cik": 1, "entityName": "TEST CO",
        "facts": {
            "us-gaap": {tag: {"units": {"USD": vals}} for tag, vals in tags.items()},
            "dei": {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
                fact(shares, "2024-02-15", "2024-02-20")]}}},
        },
    }


def year(fy: int, val: float, filed: str | None = None, **kw) -> dict:
    """A full fiscal-year duration fact ending Dec 31."""
    return fact(val, f"{fy}-12-31", filed or f"{fy + 1}-02-15", start=f"{fy}-01-01", **kw)


# ---------------------------------------------------------------- real data (trimmed AEHR companyfacts)

def test_aehr_real_values_and_tag_choice(aehr_facts):
    fin = extract_financials(aehr_facts, "AEHR")
    assert list(fin.table.columns) == [2022, 2023, 2024, 2025, 2026]
    # Values from AEHR's FY2024 10-K (fiscal year ended May 31, 2024).
    assert fin.table.loc["Revenue", 2024] == pytest.approx(66_221_000, rel=1e-3)
    assert fin.table.loc["Net Income", 2024] == pytest.approx(33_200_000, rel=1e-2)
    assert fin.sources["Revenue"] == "RevenueFromContractWithCustomerExcludingAssessedTax"
    # AEHR does not report DepreciationDepletionAndAmortization -> second fallback is used and noted.
    assert fin.sources["D&A"] == "DepreciationAndAmortization"
    assert any(n.category == "Fallback" and n.item == "D&A" for n in fin.notes)


def test_aehr_quarterly_value_inside_10k_is_excluded(aehr_facts):
    # AEHR's 10-K tags a quarter ending 2022-02-28 under Revenues with fp=FY; the duration filter drops it.
    revenues = annual_facts(aehr_facts, "Revenues", duration=True)
    assert all(d.month == 5 for d in revenues.index)


# ---------------------------------------------------------------- synthetic edge cases

def test_restatement_keeps_latest_filing():
    tag = "RevenueFromContractWithCustomerExcludingAssessedTax"
    facts = facts_with({tag: [
        year(2022, 100.0, filed="2023-02-15"),
        year(2022, 110.0, filed="2024-02-15"),     # restated in the next year's 10-K -> wins
        year(2023, 120.0, filed="2024-02-15"),
    ]})
    s = annual_facts(facts, tag, duration=True)
    assert s.loc["2022-12-31"] == 110.0
    assert len(s) == 2


def test_non_annual_facts_are_ignored():
    tag = "RevenueFromContractWithCustomerExcludingAssessedTax"
    facts = facts_with({tag: [
        year(2023, 120.0),
        fact(30.0, "2023-09-30", "2023-11-01", form="10-Q", fp="Q3", start="2023-07-01"),   # 10-Q
        fact(35.0, "2023-12-31", "2024-02-15", start="2023-10-01"),                          # Q4 inside the 10-K
        fact(999.0, "2023-12-31", "2024-05-01", form="10-K/A", start="2023-01-01"),          # not form 10-K
    ]})
    s = annual_facts(facts, tag, duration=True)
    assert s.tolist() == [120.0]


def test_fallback_is_resolved_per_year():
    """Company switches from SalesRevenueNet (old) to the ASC 606 tag: each year uses whichever exists."""
    facts = facts_with({
        "SalesRevenueNet": [year(2019, 90.0), year(2020, 95.0)],
        "RevenueFromContractWithCustomerExcludingAssessedTax": [year(2021, 100.0), year(2022, 105.0),
                                                                year(2023, 110.0)],
        "NetIncomeLoss": [year(y, 10.0) for y in range(2019, 2024)],
    })
    fin = extract_financials(facts, "TEST")
    assert fin.table.loc["Revenue"].tolist() == [90.0, 95.0, 100.0, 105.0, 110.0]
    assert "SalesRevenueNet (FY2019, FY2020)" in fin.sources["Revenue"]
    assert any(n.category == "Fallback" and n.item == "Revenue" for n in fin.notes)


def test_missing_tag_is_blank_plus_note_not_exception():
    facts = facts_with({
        "RevenueFromContractWithCustomerExcludingAssessedTax": [year(y, 100.0 + y - 2019) for y in range(2019, 2024)],
        "NetIncomeLoss": [year(y, 10.0) for y in range(2019, 2024)],
    })   # no D&A, cash flow, cash, debt, tax tags at all
    fin = extract_financials(facts, "TEST")
    assert all(math.isnan(v) for v in fin.table.loc["D&A"])
    assert any(n.category == "Missing" and n.item == "D&A" for n in fin.notes)
    assert fin.table.loc["Debt"].tolist() == [0.0] * 5
    assert any(n.category == "Assumption" and n.item == "Debt" for n in fin.notes)


def test_missing_tag_is_blank_in_workbook_and_listed_in_data_notes(aehr_model, tmp_path):
    """Knock out a line item from the real model, write the workbook, and check the sheet output."""
    model = copy.deepcopy(aehr_model)
    fin = model.fin
    fin.table.loc["CapEx"] = float("nan")
    fin.notes.append(DataNote("Missing", "CapEx", "No value under any of PaymentsToAcquirePropertyPlantAndEquipment"))
    path = write_workbook(model, tmp_path / "missing.xlsx")

    wb = load_workbook(path)
    hist = wb["Historicals"]
    capex_row = next(r for r in range(1, 30) if hist[f"A{r}"].value == "CapEx")
    assert all(hist.cell(row=capex_row, column=c).value is None for c in range(2, 7))
    notes = [(row[0], row[1]) for row in wb["Data Notes"].iter_rows(values_only=True)]
    assert ("Missing", "CapEx") in notes
