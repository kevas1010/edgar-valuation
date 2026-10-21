"""DCF math, two ways:

1. The Python reference DCF against a small example worked by hand.
2. The actual workbook formulas, evaluated by the `formulas` engine (a pure-Python Excel
   calculator, so no Excel or LibreOffice needed), against the Python reference within $0.01.
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from model.dcf import dcf, wacc
from model.pipeline import python_values

# ---------------------------------------------------------------- 1. hand-computed example


def test_dcf_matches_hand_calculation():
    """
    Base revenue 100, growth 10%/yr for 2 years, EBITDA margin 30%, D&A = CapEx = 5% of revenue,
    tax 25%, discount rate 10%, terminal growth 2%, debt 20, cash 10, 10 shares.

    Year 1: revenue 110.0, EBIT = 33.0 − 5.5 = 27.5, NOPAT = 27.5 × 0.75 = 20.625
            FCF = 20.625 + 5.5 − 5.5 = 20.625, PV = 20.625 / 1.1  = 18.75
    Year 2: revenue 121.0, EBIT = 36.3 − 6.05 = 30.25, NOPAT = 22.6875
            FCF = 22.6875, PV = 22.6875 / 1.21 = 18.75
    Terminal value = 22.6875 × 1.02 / (0.10 − 0.02) = 289.265625, PV = 289.265625 / 1.21 = 239.0625
    EV = 18.75 + 18.75 + 239.0625 = 276.5625 → equity = 276.5625 − 20 + 10 = 266.5625 → $26.65625 / share
    """
    r = dcf(base_revenue=100, growth=[0.10, 0.10], ebitda_margin=0.30, da_pct=0.05, capex_pct=0.05,
            tax_rate=0.25, discount_rate=0.10, terminal_growth=0.02, debt=20, cash=10, shares=10)
    assert r.ufcf == pytest.approx([20.625, 22.6875])
    assert r.pv_ufcf == pytest.approx([18.75, 18.75])
    assert r.terminal_value == pytest.approx(289.265625)
    assert r.enterprise_value == pytest.approx(276.5625)
    assert r.value_per_share == pytest.approx(26.65625)


def test_no_tax_on_losses_and_terminal_capex_switch():
    # EBIT = 100 × (5% − 10%) = −5 → no tax. CapEx (20) ≠ D&A (10), so the switch matters.
    kw = dict(base_revenue=100, growth=[0.0], ebitda_margin=0.05, da_pct=0.10, capex_pct=0.20, tax_rate=0.25,
              discount_rate=0.10, terminal_growth=0.0, debt=0, cash=0, shares=1)
    on, off = dcf(**kw, terminal_capex_equals_da=True), dcf(**kw, terminal_capex_equals_da=False)
    assert on.ufcf == pytest.approx([-5 + 10 - 20])       # projection year always uses CapEx % of revenue
    assert on.terminal_fcf == pytest.approx(-5.0)         # NOPAT + D&A − D&A
    assert off.terminal_fcf == pytest.approx(-15.0)       # NOPAT + D&A − CapEx


def test_wacc_ignores_immaterial_debt():
    ke = 0.04 + 1.2 * 0.05                                                    # CAPM = 10%
    assert wacc(0.04, 1.2, 0.05, 0.06, 0.25, market_cap=1000, debt=10, debt_materiality=0.05) == pytest.approx(ke)
    # 20% debt: 0.8 × 10% + 0.2 × 6% × (1 − 25%) = 8.9%
    assert wacc(0.04, 1.2, 0.05, 0.06, 0.25, market_cap=800, debt=200, debt_materiality=0.05) == pytest.approx(0.089)


# ---------------------------------------------------------------- 2. the real workbook's formulas

class Workbook:
    """Evaluate a written .xlsx with the `formulas` engine and look cells up by address or Excel name."""

    def __init__(self, path):
        import formulas
        self.path = path
        self.model = formulas.ExcelModel().loads(str(path)).finish()
        self.names = {n: dn.attr_text for n, dn in load_workbook(path).defined_names.items()}

    def key(self, ref: str) -> str:
        """'DCF!B50' or a defined name like 'DCFValue' -> the engine's cell key."""
        ref = self.names.get(ref, ref).replace("$", "")
        sheet, cell = ref.rsplit("!", 1)
        return f"'[{self.path.name}]{sheet.strip(chr(39)).upper()}'!{cell}"

    def calculate(self, inputs: dict[str, float] | None = None) -> dict:
        sol = self.model.calculate(inputs={self.key(k): v for k, v in (inputs or {}).items()})
        return {k: v.value[0][0] if hasattr(v, "value") else v for k, v in sol.items()}

    def get(self, sol: dict, ref: str):
        return sol[self.key(ref)]


@pytest.fixture(scope="module")
def evaluated(aehr_workbook):
    wb = Workbook(aehr_workbook)
    return wb, wb.calculate()


# Both sides are exact double-precision arithmetic, so the match is held far tighter than the
# spec's $0.01. A loose tolerance would hide real bugs: AEHR's value is mostly cash, so a +1%
# error in the terminal value moves value per share by only ~$0.002.
EXACT = dict(rel=1e-9)


def test_workbook_formulas_match_python_reference(evaluated, aehr_model):
    wb, sol = evaluated
    py = python_values(aehr_model)
    assert wb.get(sol, "WACC") == pytest.approx(py["WACC"], **EXACT)
    assert wb.get(sol, "EnterpriseValue") == pytest.approx(py["BaseEV"], **EXACT)
    assert wb.get(sol, "DCFValue") == pytest.approx(py["Base"], **EXACT)
    assert abs(wb.get(sol, "DCFValue") - py["Base"]) < 0.01                       # the spec's requirement
    assert wb.get(sol, "ScenarioTarget") == pytest.approx(py["Target"], **EXACT)
    assert wb.get(sol, "CompsValue") == pytest.approx(py["Comps"], **EXACT)
    assert wb.get(sol, "Summary!C9") == pytest.approx(py["Target"], **EXACT)
    for name, col in (("Bear", "B"), ("Base", "C"), ("Bull", "D")):
        assert wb.get(sol, f"Scenarios!{col}15") == pytest.approx(py[name], **EXACT)


def test_sensitivity_center_equals_dcf_value(evaluated):
    wb, sol = evaluated
    assert wb.get(sol, "Sensitivity!E10") == pytest.approx(wb.get(sol, "DCFValue"), abs=1e-9)
    assert wb.get(sol, "ModelCheck") == "OK"
    assert wb.get(sol, "Scenarios!B20") == "OK"


def test_changing_terminal_growth_updates_summary_target(evaluated, aehr_model):
    wb, sol = evaluated
    bumped = wb.calculate(inputs={"TerminalGrowth": 0.04})
    before, after = wb.get(sol, "Summary!C9"), wb.get(bumped, "Summary!C9")
    assert after != pytest.approx(before, abs=1e-6)
    assert after == pytest.approx(python_values(aehr_model, terminal_growth=0.04)["Target"], **EXACT)


def test_no_excel_error_values_anywhere(evaluated):
    wb, sol = evaluated
    errors = {k: v for k, v in sol.items() if str(v).startswith("#")}
    assert not errors, f"Formula errors: {list(errors.items())[:10]}"


def test_outputs_are_live_formulas_not_pasted_values(aehr_workbook):
    wb = load_workbook(aehr_workbook)
    for ref in ("C6", "C7", "C8", "C9", "D9", "C12", "C13", "C14"):
        assert str(wb["Summary"][ref].value).startswith("="), f"Summary!{ref} is not a formula"
    dcf_sheet, sens = wb["DCF"], wb["Sensitivity"]
    assert all(str(dcf_sheet[f"{c}16"].value).startswith("=") for c in "CDEFG")      # unlevered FCF row
    grid = [sens.cell(row=r, column=c).value for r in range(6, 15) for c in range(3, 8)]
    assert len(grid) == 45 and all(str(v).startswith("=") for v in grid)             # 9 WACC × 5 growth
    assert sens.conditional_formatting, "Sensitivity grid should have a color scale"
